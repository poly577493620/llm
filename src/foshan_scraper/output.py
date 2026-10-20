"""导出 Excel 结果."""

import json
from pathlib import Path
from typing import Any

import structlog
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from foshan_scraper.models import LlmMatchResult

log = structlog.get_logger()

# Excel 列定义
EXCEL_COLUMNS = [
    ("核验状态", "verification_status"),
    ("row_guid", "row_guid"),
    ("详情链接", "detail_url"),
    ("模型名称", "model_name"),
    # 模型识别值
    ("模型_成交时间", "llm_transaction_date"),
    ("模型_竞得人", "llm_winner"),
    ("模型_交易面积", "llm_transaction_area"),
    ("模型_成交地价", "llm_transaction_price"),
    ("模型_地块位置", "llm_land_location"),
    ("模型_土地实际用途", "llm_actual_land_usage"),
    ("模型_不确定字段", "llm_uncertain_fields"),
    ("模型_备注", "llm_notes"),
    # 参考值
    ("参考_成交时间", "ref_transaction_date"),
    ("参考_竞得人", "ref_winner"),
    ("参考_交易面积", "ref_transaction_area"),
    ("参考_成交地价", "ref_transaction_price"),
    ("参考_地块位置", "ref_land_location"),
    ("参考_土地用途分类", "ref_land_usage_category"),
    # 土地用途口径差异
    ("土地用途核验结论", "land_usage_verification"),
    # 比较结果
    ("匹配_成交时间", "match_transaction_date"),
    ("匹配_竞得人", "match_winner"),
    ("匹配_交易面积", "match_transaction_area"),
    ("匹配_成交地价", "match_transaction_price"),
    ("匹配_地块位置", "match_land_location"),
    ("匹配_土地用途", "match_land_usage"),
    ("普通差异字段", "mismatch_fields"),
    ("字段口径差异", "semantic_difference_fields"),
    ("差异说明", "mismatch_notes"),
    ("解析状态", "parse_status"),
]


def export_to_excel(results: list[LlmMatchResult], output_path: Path) -> None:
    """将比对结果导出到 Excel.

    Args:
        results: LlmMatchResult 列表
        output_path: 输出文件路径
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    if ws is None:
        msg = "Failed to create active worksheet"
        raise RuntimeError(msg)
    ws.title = "大模型识别核验结果"

    # 写入表头
    headers = [col[0] for col in EXCEL_COLUMNS]
    for col_idx, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_idx, value=header)

    # 写入数据
    for row_idx, result in enumerate(results, 2):
        ref_dump = result.reference.model_dump()
        llm_dump = result.llm_result.model_dump()

        row_data: dict[str, Any] = {
            "verification_status": result.verification_status,
            "row_guid": result.row_guid,
            "detail_url": result.detail_url,
            "model_name": result.model_name,
            # 模型识别值
            "llm_transaction_date": llm_dump.get("transaction_date"),
            "llm_winner": llm_dump.get("winner"),
            "llm_transaction_area": llm_dump.get("transaction_area"),
            "llm_transaction_price": llm_dump.get("transaction_price"),
            "llm_land_location": llm_dump.get("land_location"),
            "llm_actual_land_usage": result.llm_actual_land_usage,
            "llm_uncertain_fields": ", ".join(llm_dump.get("uncertain_fields") or []),
            "llm_notes": llm_dump.get("notes"),
            # 参考值
            "ref_transaction_date": ref_dump.get("transaction_date"),
            "ref_winner": ref_dump.get("winner"),
            "ref_transaction_area": ref_dump.get("transaction_area"),
            "ref_transaction_price": ref_dump.get("transaction_price"),
            "ref_land_location": ref_dump.get("land_location"),
            "ref_land_usage_category": result.reference_land_usage_category,
            # 土地用途口径差异
            "land_usage_verification": result.land_usage_verification,
            # 比较结果
            "match_transaction_date": "✓" if result.transaction_date_match else "✗",
            "match_winner": "✓" if result.winner_match else "✗",
            "match_transaction_area": "✓" if result.transaction_area_match else "✗",
            "match_transaction_price": "✓" if result.transaction_price_match else "✗",
            "match_land_location": "✓" if result.land_location_match else "✗",
            "match_land_usage": "✓" if result.land_usage_match else "✗",
            "mismatch_fields": ", ".join(result.mismatch_fields),
            "semantic_difference_fields": ", ".join(result.semantic_difference_fields),
            "mismatch_notes": result.mismatch_notes,
            "parse_status": result.parse_status,
        }

        for col_idx, (_, key) in enumerate(EXCEL_COLUMNS, 1):
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ""))

    # 自动调整列宽
    for col_idx in range(1, len(EXCEL_COLUMNS) + 1):
        col_letter = get_column_letter(col_idx)
        header_text = headers[col_idx - 1]
        max_length = len(header_text)
        for row_idx in range(2, len(results) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value:
                max_length = max(max_length, len(str(cell_value)))
        ws.column_dimensions[col_letter].width = min(max_length + 2, 60)

    # 冻结首行
    ws.freeze_panes = "A2"

    wb.save(output_path)
    log.info("excel_exported", path=str(output_path), count=len(results))


def save_raw_output(raw_response: str, index: int, output_dir: Path) -> Path:
    """保存模型原始响应到 JSON 文件.

    Args:
        raw_response: 模型原始输出
        index: 记录序号（从1开始）
        output_dir: raw_outputs 目录

    Returns:
        保存的文件路径
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"raw_{index:02d}.json"

    with path.open("w", encoding="utf-8") as f:
        json.dump({"index": index, "raw_response": raw_response}, f, ensure_ascii=False, indent=2)
    return path
