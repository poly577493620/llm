"""测试 Excel 导出模块（不访问真实网络）."""

from pathlib import Path

from openpyxl import load_workbook

from foshan_scraper.models import LlmMatchResult, LlmRecognitionResult, ReferenceRecord
from foshan_scraper.output import export_to_excel


def _make_match_result(
    guid: str = "test-guid-001",
    status: str = "match",
    winner_match: bool = True,
    land_usage_verification: str = "完全匹配",
    semantic_difference_fields: list[str] | None = None,
    mismatch_fields: list[str] | None = None,
) -> LlmMatchResult:
    """创建测试用 LlmMatchResult."""
    ref = ReferenceRecord(
        row_guid=guid,
        transaction_date="2026/7/27",
        winner="参考公司",
        transaction_area="4840.75",
        transaction_price="地价:15746万元",
        land_location="佛山市禅城区",
        land_usage="商业服务业用地·商业用地",
        detail_url="http://example.com/detail",
    )
    llm = LlmRecognitionResult(
        transaction_date="2026/7/27" if status == "match" else "2026/7/28",
        winner="参考公司" if winner_match else "错误公司",
        transaction_area="4840.75",
        transaction_price="地价:15746万元",
        land_location="佛山市禅城区",
        land_usage="商业用地（零售商业用地、餐饮用地、旅馆用地），兼容娱乐用地",
    )
    return LlmMatchResult(
        row_guid=guid,
        detail_url="http://example.com/detail",
        reference=ref,
        llm_result=llm,
        model_name="gpt-4o",
        raw_response="{}",
        transaction_date_match=status == "match",
        winner_match=winner_match,
        transaction_area_match=True,
        transaction_price_match=True,
        land_location_match=True,
        land_usage_match=land_usage_verification == "完全匹配",
        verification_status=status,
        mismatch_fields=mismatch_fields if mismatch_fields is not None else ([] if status == "match" else ["winner"]),
        semantic_difference_fields=semantic_difference_fields or [],
        mismatch_notes="竞得人不匹配" if status != "match" else "",
        reference_land_usage_category="商业服务业用地·商业用地",
        llm_actual_land_usage="商业用地（零售商业用地、餐饮用地、旅馆用地），兼容娱乐用地",
        land_usage_verification=land_usage_verification,
    )


def test_export_to_excel(tmp_path: Path) -> None:
    """测试 Excel 导出（mock 数据）."""
    results = [_make_match_result("guid-1", "match"), _make_match_result("guid-2", "needs_review")]
    output_path = tmp_path / "test_output.xlsx"

    export_to_excel(results, output_path)

    assert output_path.exists()
    assert output_path.stat().st_size > 0


def test_export_creates_parent_dirs(tmp_path: Path) -> None:
    """测试自动创建父目录."""
    results = [_make_match_result()]
    output_path = tmp_path / "nested" / "dir" / "test.xlsx"

    export_to_excel(results, output_path)

    assert output_path.exists()


def test_excel_land_usage_columns(tmp_path: Path) -> None:
    """Excel 正确写入土地用途口径差异专用列."""
    results = [
        _make_match_result(
            "guid-semantic",
            "needs_review",
            winner_match=True,
            land_usage_verification="字段口径不同，需人工核验",
            semantic_difference_fields=["land_usage"],
            mismatch_fields=[],
        ),
    ]
    output_path = tmp_path / "test_land_usage.xlsx"
    export_to_excel(results, output_path)

    wb = load_workbook(output_path)
    ws = wb.active
    assert ws is not None

    # 找到列索引
    headers = [cell.value for cell in ws[1]]
    idx_ref_category = headers.index("参考_土地用途分类") + 1
    idx_llm_actual = headers.index("模型_土地实际用途") + 1
    idx_verification = headers.index("土地用途核验结论") + 1
    idx_semantic = headers.index("字段口径差异") + 1
    idx_mismatch = headers.index("普通差异字段") + 1

    row2 = [cell.value for cell in ws[2]]
    assert row2[idx_ref_category - 1] == "商业服务业用地·商业用地"
    assert row2[idx_llm_actual - 1] == "商业用地（零售商业用地、餐饮用地、旅馆用地），兼容娱乐用地"
    assert row2[idx_verification - 1] == "字段口径不同，需人工核验"
    assert row2[idx_semantic - 1] == "land_usage"
    assert row2[idx_mismatch - 1] is None or row2[idx_mismatch - 1] == ""


def test_excel_all_match_row(tmp_path: Path) -> None:
    """全部匹配时口径差异列为空."""
    results = [
        _make_match_result(
            "guid-match",
            "match",
            land_usage_verification="完全匹配",
        ),
    ]
    output_path = tmp_path / "test_match.xlsx"
    export_to_excel(results, output_path)

    wb = load_workbook(output_path)
    ws = wb.active
    assert ws is not None

    headers = [cell.value for cell in ws[1]]
    idx_verification = headers.index("土地用途核验结论") + 1
    idx_semantic = headers.index("字段口径差异") + 1
    idx_mismatch = headers.index("普通差异字段") + 1

    row2 = [cell.value for cell in ws[2]]
    assert row2[idx_verification - 1] == "完全匹配"
    assert row2[idx_semantic - 1] is None or row2[idx_semantic - 1] == ""
    assert row2[idx_mismatch - 1] is None or row2[idx_mismatch - 1] == ""


def test_excel_real_mismatch_plus_semantic(tmp_path: Path) -> None:
    """真正差异和口径差异同时存在时，两列都有值."""
    results = [
        _make_match_result(
            "guid-both",
            "needs_review",
            winner_match=False,
            land_usage_verification="字段口径不同，需人工核验",
            semantic_difference_fields=["land_usage"],
            mismatch_fields=["winner"],
        ),
    ]
    output_path = tmp_path / "test_both.xlsx"
    export_to_excel(results, output_path)

    wb = load_workbook(output_path)
    ws = wb.active
    assert ws is not None

    headers = [cell.value for cell in ws[1]]
    idx_semantic = headers.index("字段口径差异") + 1
    idx_mismatch = headers.index("普通差异字段") + 1

    row2 = [cell.value for cell in ws[2]]
    assert row2[idx_semantic - 1] == "land_usage"
    assert row2[idx_mismatch - 1] == "winner"
